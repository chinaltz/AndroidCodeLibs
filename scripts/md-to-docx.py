#!/usr/bin/env python3
"""
Markdown → Word 文档（.docx）
用法:
  python3 scripts/md-to-docx.py content/posts/ai-for-ordinary-people
  python3 scripts/md-to-docx.py --all
"""

import sys, os, re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent.parent
POSTS_DIR = ROOT / 'content' / 'posts'
BOX_CORNER = re.compile(r'[╔╗╚╝┌┐└┘│├┤┬┴┼╠╣╦╩╤╧═]')
FONT_BODY = 'Microsoft YaHei'
FONT_CODE = 'Consolas'

def set_cell_border(cell, **kwargs):
    """设置单元格边框"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('start', 'top', 'end', 'bottom', 'insideH', 'insideV'):
        if edge in kwargs:
            el = OxmlElement(f'w:{edge}')
            for attr, val in kwargs[edge].items():
                el.set(qn(f'w:{attr}'), str(val))
            tcBorders.append(el)
    tcPr.append(tcBorders)

def extract_title(md_text):
    m = re.search(r'^#\s+(.+)$', md_text, re.MULTILINE)
    return m.group(1).strip() if m else '文章'

def is_table_sep(line):
    return bool(re.match(r'^\|[\s\-:|]+\|$', line.strip()))

def parse_table_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]

def split_inline(text):
    """拆分内联 Markdown 元素"""
    parts = []
    regex = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|!\[[^\]]*\]\([^)]+\)|\[[^\]]+\]\([^)]+\))')
    segments = regex.split(text)
    for seg in segments:
        if not seg:
            continue
        m = re.match(r'^\*\*(.+)\*\*$', seg)
        if m:
            parts.append(('bold', m.group(1)))
            continue
        m = re.match(r'^\*(.+)\*$', seg)
        if m:
            parts.append(('italic', m.group(1)))
            continue
        m = re.match(r'^`(.+)`$', seg)
        if m:
            parts.append(('code', m.group(1)))
            continue
        m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)$', seg)
        if m:
            parts.append(('image', m.group(1), m.group(2)))
            continue
        m = re.match(r'^\[([^\]]+)\]\(([^)]+)\)$', seg)
        if m:
            parts.append(('link', m.group(1), m.group(2)))
            continue
        parts.append(('text', seg))
    return parts

def add_inline_run(paragraph, elem):
    """向段落添加内联元素 run"""
    if elem[0] == 'bold':
        run = paragraph.add_run(elem[1])
        run.bold = True
    elif elem[0] == 'italic':
        run = paragraph.add_run(elem[1])
        run.italic = True
    elif elem[0] == 'code':
        run = paragraph.add_run(elem[1])
        run.font.name = FONT_CODE
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x50, 0x50, 0x50)
    elif elem[0] == 'link':
        run = paragraph.add_run(elem[1])
        run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        run.underline = True
    elif elem[0] == 'text':
        run = paragraph.add_run(elem[1])
    else:
        return
    return run

def add_body_paragraph(doc, text):
    """普通正文段落"""
    p = doc.add_paragraph()
    for elem in split_inline(text):
        add_inline_run(p, elem)
    return p

def add_code_block(doc, code_text):
    """代码块：缩进 + Consolas + 灰色文字"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(1)
    run = p.add_run(code_text.strip())
    run.font.name = FONT_CODE
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x50, 0x50, 0x50)
    return p

def add_quote_block(doc, text):
    """引用块：缩进 + 斜体 + 灰色"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(1)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x64, 0x64, 0x64)
    return p

def add_separator(doc):
    """分隔线"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('─' * 50)
    return p

def add_table_from_md(doc, rows):
    """从 Markdown 表格行列表创建 Word 表格"""
    if len(rows) < 2:
        return
    header = rows[0]
    body = rows[1:]
    num_cols = len(header)
    table = doc.add_table(rows=1 + len(body), cols=num_cols)
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 表头
    for j, cell_text in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(cell_text)
        run.bold = True
        run.font.size = Pt(9)

    # 数据行
    for i, row_data in enumerate(body):
        for j, cell_text in enumerate(row_data):
            if j >= num_cols:
                break
            cell = table.rows[i + 1].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            for elem in split_inline(cell_text):
                run = add_inline_run(p, elem)
                if run:
                    run.font.size = Pt(10)

    doc.add_paragraph()  # 表后空行

def add_image(doc, img_path, post_dir, alt=''):
    """嵌入图片，宽度 14cm"""
    full_path = (Path(post_dir) / img_path.lstrip('./')).resolve()
    if full_path.exists():
        try:
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(str(full_path), width=Cm(14))
            return p
        except Exception:
            pass
    # 图片不存在时显示占位文字
    return doc.add_paragraph(f'[图片: {alt or img_path}]')

def is_flow_diagram(code):
    """含箭头/emoji 但无框线四角"""
    return bool(re.search(r'[─→↓]', code)) and bool(re.search(r'[📋✍️🎥✂️📤✅🙋❶❷❸①②③🎼🎤]', code)) and not BOX_CORNER.search(code)

def preprocess_ascii_boxes(md_text):
    """将未包在 ``` 里的 ASCII 框线图用 ``` 包裹"""
    lines = md_text.split('\n')
    out = []
    buf = []
    in_box = False
    in_fence = False

    for line in lines:
        if line.strip().startswith('```'):
            if in_box:
                out.append('```')
                out.extend(buf)
                out.append('```')
                buf = []
                in_box = False
            out.append(line)
            in_fence = not in_fence
            continue
        if in_fence:
            out.append(line)
            continue
        stripped = line.strip()
        is_box = len(stripped) > 0 and BOX_CORNER.search(line) and not line.startswith('|')
        if is_box:
            in_box = True
            buf.append(line)
        else:
            if in_box:
                out.append('```')
                out.extend(buf)
                out.append('```')
                buf = []
                in_box = False
            out.append(line)
    if in_box:
        out.append('```')
        out.extend(buf)
        out.append('```')
    return '\n'.join(out)

def markdown_to_docx(md_text, post_dir):
    """将 Markdown 文本转换为 Word 文档"""
    doc = Document()

    # 默认字体
    style = doc.styles['Normal']
    style.font.name = FONT_BODY
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # 修改标题样式字体
    for level in range(1, 5):
        heading_style = doc.styles[f'Heading {level}']
        heading_style.font.name = FONT_BODY
        heading_style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    lines = md_text.split('\n')
    i = 0

    while i < len(lines):
        line = lines[i]

        # 空行
        if line.strip() == '':
            i += 1
            continue

        # 代码围栏
        if line.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code_lines.append(lines[i])
                i += 1
            code_text = '\n'.join(code_lines)
            if BOX_CORNER.search(code_text):
                # ASCII 框线图
                for cl in code_lines:
                    p = doc.add_paragraph()
                    run = p.add_run(cl)
                    run.font.name = FONT_CODE
                    run.font.size = Pt(7.5)
                    run.font.color.rgb = RGBColor(0x3C, 0x3C, 0x3C)
            elif is_flow_diagram(code_text):
                # 流程图
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(code_text.strip())
                run.font.name = FONT_CODE
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x50, 0x50, 0x50)
            else:
                # 普通代码块
                for cl in code_lines:
                    p = doc.add_paragraph()
                    pf = p.paragraph_format
                    pf.left_indent = Cm(1)
                    run = p.add_run(cl)
                    run.font.name = FONT_CODE
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x50, 0x50, 0x50)
            i += 1
            continue

        # 分隔线
        if re.match(r'^---+$', line.strip()):
            add_separator(doc)
            i += 1
            continue

        # 表格
        if line.startswith('|') and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            rows = [parse_table_row(line)]
            i += 2
            while i < len(lines) and lines[i].startswith('|'):
                rows.append(parse_table_row(lines[i]))
                i += 1
            add_table_from_md(doc, rows)
            continue

        # 引用块
        if line.startswith('>'):
            bq_lines = []
            while i < len(lines) and (lines[i].startswith('>') or (lines[i].strip() == '' and i + 1 < len(lines) and lines[i + 1].startswith('>'))):
                if lines[i].strip():
                    bq_lines.append(lines[i])
                i += 1
            bq_text = ' '.join(l.replace('> ', '', 1).replace('>', '', 1) for l in bq_lines)
            add_quote_block(doc, bq_text)
            continue

        # 标题
        m = re.match(r'^####\s+(.+)$', line)
        if m:
            p = doc.add_heading(m.group(1), level=4)
            i += 1
            continue

        m = re.match(r'^###\s+(.+)$', line)
        if m:
            p = doc.add_heading(m.group(1), level=3)
            i += 1
            continue

        m = re.match(r'^##\s+(.+)$', line)
        if m:
            p = doc.add_heading(m.group(1), level=2)
            i += 1
            continue

        m = re.match(r'^#\s+(.+)$', line)
        if m:
            p = doc.add_heading(m.group(1), level=1)
            i += 1
            continue

        # 单独图片行
        m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', line)
        if m:
            add_image(doc, m.group(2), post_dir, m.group(1))
            i += 1
            continue

        # 无序列表
        if re.match(r'^[-*]\s+', line):
            items = []
            while i < len(lines) and re.match(r'^[-*]\s+', lines[i]):
                items.append(re.sub(r'^[-*]\s+', '', lines[i]))
                i += 1
            for item in items:
                p = doc.add_paragraph(style='List Bullet')
                for elem in split_inline(item):
                    add_inline_run(p, elem)
            continue

        # 有序列表
        if re.match(r'^\d+\.\s+', line):
            items = []
            while i < len(lines) and re.match(r'^\d+\.\s+', lines[i]):
                items.append(re.sub(r'^\d+\.\s+', '', lines[i]))
                i += 1
            for item in items:
                p = doc.add_paragraph(style='List Number')
                for elem in split_inline(item):
                    add_inline_run(p, elem)
            continue

        # 普通段落
        para_lines = []
        while (i < len(lines) and lines[i].strip() != '' and
               not lines[i].startswith('#') and
               not lines[i].startswith('>') and
               not lines[i].startswith('```') and
               not lines[i].startswith('|') and
               not re.match(r'^---+$', lines[i].strip()) and
               not re.match(r'^[-*]\s+', lines[i]) and
               not re.match(r'^\d+\.\s+', lines[i]) and
               not re.match(r'^!\[', lines[i])):
            para_lines.append(lines[i])
            i += 1
        if para_lines:
            text = ' '.join(para_lines)
            # 如果整段是框线图（未被围栏包裹的情况）
            if all(BOX_CORNER.search(pl) or pl.strip() == '' for pl in para_lines):
                for pl in para_lines:
                    p = doc.add_paragraph()
                    run = p.add_run(pl)
                    run.font.name = FONT_CODE
                    run.font.size = Pt(7.5)
                    run.font.color.rgb = RGBColor(0x3C, 0x3C, 0x3C)
            else:
                add_body_paragraph(doc, text)
        else:
            add_body_paragraph(doc, line)
            i += 1

    return doc

def convert_post(post_dir):
    md_path = Path(post_dir) / 'index.md'
    if not md_path.exists():
        print(f'跳过（无 index.md）: {post_dir}')
        return False

    md_raw = md_path.read_text(encoding='utf-8')
    md_text = preprocess_ascii_boxes(md_raw)
    title = extract_title(md_raw)
    doc = markdown_to_docx(md_text, post_dir)

    # 清理标题中的特殊字符用于文件名
    safe_title = title.replace('/', '-').replace('\\', '-').replace(':', '-').replace('"', "'")
    out_path = Path(post_dir) / f'{safe_title}.docx'
    doc.save(str(out_path))
    print(f'✅ {out_path.relative_to(ROOT)}')
    return True

def find_all_posts():
    if not POSTS_DIR.exists():
        return []
    return [POSTS_DIR / d for d in os.listdir(POSTS_DIR)
            if (POSTS_DIR / d).is_dir() and ((POSTS_DIR / d) / 'index.md').exists()]

def resolve_input(arg):
    if not arg:
        return None
    p = (ROOT / arg).resolve()
    if p.is_file() and p.suffix == '.md':
        return str(p.parent)
    if (p / 'index.md').exists():
        return str(p)
    return None

def print_usage():
    print('''
技趣星球 · Markdown → Word 文档（.docx）

用法:
  python3 scripts/md-to-docx.py <文章目录>
  python3 scripts/md-to-docx.py --all

示例:
  python3 scripts/md-to-docx.py content/posts/ai-for-ordinary-people

输出文件: <文章目录>/<文章标题>.docx
''')

if __name__ == '__main__':
    args = sys.argv[1:]
    if '--help' in args or '-h' in args:
        print_usage()
        sys.exit(0)

    if '--all' in args:
        posts = find_all_posts()
        if not posts:
            print('未找到文章')
            sys.exit(1)
        n = 0
        for p in posts:
            if convert_post(p):
                n += 1
        print(f'\n共生成 {n} 篇')
        sys.exit(0)

    if not args:
        print_usage()
        sys.exit(1)

    post_dir = resolve_input(args[0])
    if not post_dir:
        print(f'找不到文章: {args[0]}')
        sys.exit(1)
    convert_post(post_dir)
